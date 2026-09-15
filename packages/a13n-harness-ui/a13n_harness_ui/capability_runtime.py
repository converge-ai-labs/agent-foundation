"""Release-owned Harness UI definition capabilities and fresh run collaborators."""

from __future__ import annotations

import ipaddress
import socket
import ssl
from collections.abc import AsyncIterable, AsyncIterator, Collection, Iterable
from dataclasses import replace
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, cast
from urllib.parse import parse_qs, urlencode, urljoin, urlsplit, urlunsplit
from uuid import uuid4

import httpcore2
import httpx2
from a13n_harness.capabilities import (
    DocumentConversionError,
    DocumentConversionRequest,
    DocumentConversionResult,
    WebBinding,
    WebPolicy,
    WebProviderError,
    WebRequest,
    WebResponse,
    WebScrapeRequest,
    WebScrapeResult,
    WebSearchRequest,
    WebSearchResponse,
    WebSearchResult,
)
from a13n_harness.capabilities.documents import DOCUMENTS_CAPABILITY_ID
from a13n_harness.capabilities.web import WEB_CAPABILITY_ID
from a13n_harness.context import AgentContext, RunBindings
from anyio import getaddrinfo, to_thread
from pydantic_ai import RunContext
from pydantic_ai.messages import FilePart

_MAX_SEARCH_RESPONSE_BYTES = 2 * 1024 * 1024
_USER_AGENT = "a13n-harness-ui/0 web tools"
_REDIRECT_STATUSES = frozenset({301, 302, 303, 307, 308})


class PublicWebPolicy(WebPolicy):
    """Allow public HTTP(S) destinations and return the exact authorized addresses."""

    async def authorize(self, url: str, *, purpose: str) -> None:
        await self.resolve(url, purpose=purpose)

    async def resolve(self, url: str, *, purpose: str) -> tuple[str, ...]:
        del purpose
        parsed = urlsplit(url)
        hostname = parsed.hostname
        if parsed.scheme not in {"http", "https"} or hostname is None:
            raise WebProviderError("web_url_invalid")
        if hostname.casefold() == "localhost" or hostname.casefold().endswith(".localhost"):
            raise WebProviderError("web_destination_denied")
        port = parsed.port or (443 if parsed.scheme == "https" else 80)
        try:
            resolved = await getaddrinfo(hostname, port, type=socket.SOCK_STREAM)
        except OSError as exc:
            raise WebProviderError("web_dns_failed") from exc
        addresses = tuple(dict.fromkeys(item[4][0] for item in resolved))
        if not addresses:
            raise WebProviderError("web_dns_failed")
        if any(not ipaddress.ip_address(address).is_global for address in addresses):
            raise WebProviderError("web_destination_denied")
        return addresses


class _PinnedNetworkBackend(httpcore2.AsyncNetworkBackend):
    """Connect original HTTP origins only through policy-authorized IP addresses."""

    def __init__(self) -> None:
        self._delegate = cast(httpcore2.AsyncNetworkBackend, httpcore2.AnyIOBackend())
        self._pins: dict[tuple[str, int], tuple[str, ...]] = {}

    def pin(self, host: str, port: int, addresses: tuple[str, ...]) -> None:
        self._pins[(host.casefold(), port)] = addresses

    async def connect_tcp(
        self,
        host: str,
        port: int,
        timeout: float | None = None,
        local_address: str | None = None,
        socket_options: Iterable[tuple[Any, ...]] | None = None,
    ) -> httpcore2.AsyncNetworkStream:
        addresses = self._pins.get((host.casefold(), port))
        if addresses is None:
            raise WebProviderError("web_destination_not_authorized")
        last_error: OSError | None = None
        for address in addresses:
            try:
                return await self._delegate.connect_tcp(
                    address,
                    port,
                    timeout=timeout,
                    local_address=local_address,
                    socket_options=socket_options,
                )
            except OSError as exc:
                last_error = exc
        raise WebProviderError("web_connect_failed") from last_error

    async def connect_unix_socket(
        self,
        path: str,
        timeout: float | None = None,
        socket_options: Iterable[tuple[Any, ...]] | None = None,
    ) -> httpcore2.AsyncNetworkStream:
        del path, timeout, socket_options
        raise WebProviderError("web_destination_denied")

    async def sleep(self, seconds: float) -> None:
        await self._delegate.sleep(seconds)


class _CoreResponseStream(httpx2.AsyncByteStream):
    def __init__(self, stream: AsyncIterable[bytes]) -> None:
        self._stream = stream

    async def __aiter__(self) -> AsyncIterator[bytes]:
        async for chunk in self._stream:
            yield chunk

    async def aclose(self) -> None:
        close = getattr(self._stream, "aclose", None)
        if close is not None:
            await close()


class _PinnedTransport(httpx2.AsyncBaseTransport):
    def __init__(self, backend: _PinnedNetworkBackend) -> None:
        self._pool = httpcore2.AsyncConnectionPool(
            ssl_context=ssl.create_default_context(),
            network_backend=backend,
        )

    async def handle_async_request(self, request: httpx2.Request) -> httpx2.Response:
        response = await self._pool.handle_async_request(
            httpcore2.Request(
                method=request.method,
                url=httpcore2.URL(
                    scheme=request.url.raw_scheme,
                    host=request.url.raw_host,
                    port=request.url.port,
                    target=request.url.raw_path,
                ),
                headers=request.headers.raw,
                content=request.stream,
                extensions=request.extensions,
            )
        )
        if not isinstance(response.stream, AsyncIterable):
            raise WebProviderError("web_response_stream_invalid")
        return httpx2.Response(
            status_code=response.status,
            headers=response.headers,
            stream=_CoreResponseStream(response.stream),
            extensions=response.extensions,
        )

    async def aclose(self) -> None:
        await self._pool.aclose()


class HttpxWebClient:
    """Bounded async HTTP transport with explicit per-hop policy checks."""

    async def request(self, request: WebRequest, *, policy: WebPolicy) -> WebResponse:
        if not isinstance(policy, PublicWebPolicy):
            raise WebProviderError("web_policy_unsupported")
        backend = _PinnedNetworkBackend()
        client = httpx2.AsyncClient(
            follow_redirects=False,
            timeout=httpx2.Timeout(request.deadline_seconds),
            transport=_PinnedTransport(backend),
            trust_env=False,
            headers={"User-Agent": _USER_AGENT},
        )
        response: httpx2.Response | None = None
        current_url = request.url
        redirects = 0
        try:
            while True:
                parsed = urlsplit(current_url)
                hostname = parsed.hostname
                if hostname is None:
                    raise WebProviderError("web_url_invalid")
                port = parsed.port or (443 if parsed.scheme == "https" else 80)
                addresses = await policy.resolve(current_url, purpose=request.purpose)
                backend.pin(hostname, port, addresses)
                response = await client.send(
                    client.build_request(request.method, current_url),
                    stream=True,
                )
                if response.status_code not in _REDIRECT_STATUSES:
                    break
                location = response.headers.get("location")
                await response.aclose()
                response = None
                if location is None:
                    raise WebProviderError("web_redirect_invalid")
                if redirects >= request.max_redirects:
                    raise WebProviderError("web_redirect_limit")
                current_url = urljoin(current_url, location)
                redirects += 1

            final_url = str(response.url)
            headers = dict(response.headers.items())

            async def body() -> AsyncIterator[bytes]:
                assert response is not None
                total = 0
                async for chunk in response.aiter_bytes(chunk_size=request.max_stream_chunk_bytes):
                    total += len(chunk)
                    if total > request.max_response_bytes:
                        raise WebProviderError("web_body_too_large")
                    yield chunk

            async def close() -> None:
                assert response is not None
                try:
                    await response.aclose()
                finally:
                    await client.aclose()

            return WebResponse(
                status_code=response.status_code,
                final_url=final_url,
                canonical_url=_canonical_url(final_url),
                headers=headers,
                body=body(),
                reason=response.reason_phrase,
                redirect_count=redirects,
                _close=close,
            )
        except BaseException:
            if response is not None:
                await response.aclose()
            await client.aclose()
            raise


class DuckDuckGoSearchProvider:
    """Keyless bounded Host search through DuckDuckGo's HTML endpoint."""

    def __init__(self, client: HttpxWebClient, policy: PublicWebPolicy) -> None:
        self._client = client
        self._policy = policy

    async def search(self, request: WebSearchRequest) -> WebSearchResponse:
        url = f"https://html.duckduckgo.com/html/?{urlencode({'q': request.query})}"
        response = await self._client.request(
            WebRequest(
                url=url,
                purpose="fetch",
                deadline_seconds=30,
                max_redirects=4,
                max_response_bytes=_MAX_SEARCH_RESPONSE_BYTES,
                max_header_count=128,
                max_header_bytes=64 * 1024,
                max_stream_chunk_bytes=64 * 1024,
            ),
            policy=self._policy,
        )
        try:
            if not 200 <= response.status_code < 300:
                raise WebProviderError("web_search_failed")
            parser = _DuckDuckGoParser(request.limit)
            parser.feed((await _read_body(response)).decode("utf-8", errors="replace"))
            return WebSearchResponse(results=tuple(parser.results))
        finally:
            await response.close()


class HtmlScrapeProvider:
    """Fetch public HTML and convert its bounded main document to Markdown."""

    supports_domain_restrictions = True

    def __init__(self, client: HttpxWebClient) -> None:
        self._client = client

    async def scrape(self, request: WebScrapeRequest, *, policy: WebPolicy) -> WebScrapeResult:
        response = await self._client.request(
            WebRequest(
                url=request.url,
                purpose="scrape",
                deadline_seconds=request.deadline_seconds,
                max_redirects=request.max_redirects,
                max_response_bytes=min(request.max_content_bytes * 4, 16 * 1024 * 1024),
                max_header_count=128,
                max_header_bytes=64 * 1024,
                max_stream_chunk_bytes=64 * 1024,
            ),
            policy=policy,
        )
        try:
            if not 200 <= response.status_code < 300:
                raise WebProviderError("web_scrape_failed")
            content_type = response.headers.get("content-type", "").casefold()
            if "html" not in content_type and "text/" not in content_type:
                raise WebProviderError("web_content_type_unsupported")
            source = (await _read_body(response)).decode("utf-8", errors="replace")
            from markdownify import markdownify

            converted = markdownify(source, heading_style="ATX").strip()
            if len(converted.encode("utf-8")) > request.max_content_bytes:
                raise WebProviderError("web_body_too_large")
            return WebScrapeResult(
                content=converted,
                source_url=response.final_url,
                canonical_url=response.canonical_url,
            )
        finally:
            await response.close()


class LocalDocumentConverter:
    """Convert common local PDF, Office, and EPUB inputs without external services."""

    async def convert(self, request: DocumentConversionRequest) -> DocumentConversionResult:
        return await to_thread.run_sync(_convert_document, request)


def production_run_bindings(
    bindings: RunBindings,
    owner_capability_ids: Collection[str],
) -> RunBindings:
    """Supply fresh Host collaborators required by one logical Harness Run."""

    if WEB_CAPABILITY_ID in owner_capability_ids:
        client = HttpxWebClient()
        policy = PublicWebPolicy()
        bindings = replace(
            bindings,
            web=WebBinding(
                client=client,
                policy=policy,
                search_provider=DuckDuckGoSearchProvider(client, policy),
                scrape_provider=HtmlScrapeProvider(client),
            ),
        )
    if DOCUMENTS_CAPABILITY_ID in owner_capability_ids:
        bindings = replace(bindings, document_converter=LocalDocumentConverter())
    return bindings


async def save_native_image(ctx: RunContext[AgentContext], image: FilePart) -> str:
    """Save model-native output in this Run's Thread scratch mount."""
    extensions = {"image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp", "image/gif": ".gif"}
    suffix = extensions.get(image.content.media_type, ".img")
    name = f"image-{uuid4().hex}{suffix}"
    environment = ctx.deps.environment
    selected = environment.select_files(name, alias="thread-files")
    assert selected.mount_path is not None
    path = f"{selected.mount_path.rstrip('/')}{selected.resolved_path.path}"

    async def chunks() -> AsyncIterator[bytes]:
        for offset in range(0, len(image.content.data), 65_536):
            yield image.content.data[offset : offset + 65_536]

    async with environment.open_files(selected) as files:
        await files.write_bytes_stream(path, chunks(), mode="create")
    return path


async def _read_body(response: WebResponse) -> bytes:
    chunks = [chunk async for chunk in response.body]
    return b"".join(chunks)


def _canonical_url(url: str) -> str:
    parsed = urlsplit(url)
    hostname = parsed.hostname or ""
    if ":" in hostname and not hostname.startswith("["):
        hostname = f"[{hostname}]"
    if parsed.port is not None:
        hostname = f"{hostname}:{parsed.port}"
    return urlunsplit((parsed.scheme, hostname, parsed.path, "", ""))


class _DuckDuckGoParser(HTMLParser):
    def __init__(self, limit: int) -> None:
        super().__init__(convert_charrefs=True)
        self._limit = limit
        self._active: str | None = None
        self._href: str | None = None
        self._title: list[str] = []
        self._snippet: list[str] = []
        self._pending: tuple[str, str] | None = None
        self.results: list[WebSearchResult] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if len(self.results) >= self._limit:
            return
        values = dict(attrs)
        classes = set((values.get("class") or "").split())
        if tag == "a" and "result__a" in classes:
            self._active = "title"
            self._href = values.get("href")
            self._title = []
        elif "result__snippet" in classes:
            self._active = "snippet"
            self._snippet = []

    def handle_data(self, data: str) -> None:
        if self._active == "title":
            self._title.append(data)
        elif self._active == "snippet":
            self._snippet.append(data)

    def handle_endtag(self, tag: str) -> None:
        if self._active == "title" and tag == "a":
            href = _search_result_url(self._href)
            title = " ".join("".join(self._title).split())
            self._pending = (title, href) if title and href is not None else None
            self._active = None
        elif self._active == "snippet" and tag in {"a", "div", "span"}:
            if self._pending is not None:
                title, url = self._pending
                snippet = " ".join("".join(self._snippet).split())
                self.results.append(WebSearchResult(title=title, url=url, snippet=snippet))
                self._pending = None
            self._active = None


def _search_result_url(value: str | None) -> str | None:
    if value is None:
        return None
    absolute = urljoin("https://duckduckgo.com", value)
    parsed = urlsplit(absolute)
    target = parse_qs(parsed.query).get("uddg", [absolute])[0]
    target_parsed = urlsplit(target)
    return target if target_parsed.scheme in {"http", "https"} and target_parsed.hostname else None


def _convert_document(request: DocumentConversionRequest) -> DocumentConversionResult:
    suffix = Path(request.source_name).suffix.casefold()
    try:
        if request.kind == "pdf":
            return _convert_pdf(request)
        if suffix == ".docx":
            import mammoth

            return _text_result(mammoth.convert_to_markdown(BytesIO(request.source_bytes)).value, request)
        if suffix == ".pptx":
            return _convert_presentation(request)
        if suffix == ".xlsx":
            return _convert_workbook(request)
        if suffix == ".epub":
            return _convert_epub(request)
    except DocumentConversionError:
        raise
    except Exception as exc:
        raise DocumentConversionError("document_conversion_failed") from exc
    raise DocumentConversionError("document_format_unsupported")


def _convert_pdf(request: DocumentConversionRequest) -> DocumentConversionResult:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(request.source_bytes))
    total = len(reader.pages)
    start = request.page_start or 1
    requested_end = total if request.page_end in (None, -1) else request.page_end
    assert requested_end is not None
    end = min(total, requested_end)
    if total == 0 or start > end:
        raise DocumentConversionError("document_page_range_invalid")
    pages = []
    for number in range(start, end + 1):
        text = reader.pages[number - 1].extract_text() or ""
        pages.append(f"## Page {number}\n\n{text.strip()}")
    markdown = "\n\n".join(pages).strip()
    _require_markdown_bound(markdown, request)
    return DocumentConversionResult(
        markdown=markdown,
        total_pages=total,
        converted_pages=end - start + 1,
        page_start=start,
        page_end=end,
    )


def _convert_presentation(request: DocumentConversionRequest) -> DocumentConversionResult:
    from pptx import Presentation
    from pptx.shapes.autoshape import Shape
    from pptx.shapes.placeholder import BasePlaceholder

    presentation = Presentation(BytesIO(request.source_bytes))
    slides: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        text = [
            shape.text.strip()
            for shape in slide.shapes
            if isinstance(shape, (Shape, BasePlaceholder)) and shape.text.strip()
        ]
        slides.append(f"## Slide {index}\n\n" + "\n\n".join(text))
    return _text_result("\n\n".join(slides), request)


def _convert_workbook(request: DocumentConversionRequest) -> DocumentConversionResult:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(request.source_bytes), read_only=True, data_only=True)
    try:
        sheets: list[str] = []
        for sheet in workbook.worksheets:
            rows = [["" if value is None else str(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            width = max((len(row) for row in rows), default=0)
            if width == 0:
                sheets.append(f"## {sheet.title}")
                continue
            normalized = [row + [""] * (width - len(row)) for row in rows]
            header = normalized[0] if normalized else [""] * width
            table = ["| " + " | ".join(_table_cell(value) for value in header) + " |"]
            table.append("| " + " | ".join("---" for _ in range(width)) + " |")
            table.extend("| " + " | ".join(_table_cell(value) for value in row) + " |" for row in normalized[1:])
            sheets.append(f"## {sheet.title}\n\n" + "\n".join(table))
        return _text_result("\n\n".join(sheets), request)
    finally:
        workbook.close()


def _convert_epub(request: DocumentConversionRequest) -> DocumentConversionResult:
    import zipfile

    from markdownify import markdownify

    sections: list[str] = []
    with TemporaryDirectory() as temporary:
        archive_path = Path(temporary) / "document.epub"
        archive_path.write_bytes(request.source_bytes)
        with zipfile.ZipFile(archive_path) as archive:
            for name in sorted(archive.namelist()):
                if Path(name).suffix.casefold() not in {".html", ".htm", ".xhtml"}:
                    continue
                source = archive.read(name).decode("utf-8", errors="replace")
                converted = markdownify(source, heading_style="ATX").strip()
                if converted:
                    sections.append(converted)
    return _text_result("\n\n".join(sections), request)


def _text_result(markdown: str, request: DocumentConversionRequest) -> DocumentConversionResult:
    value = markdown.strip()
    _require_markdown_bound(value, request)
    return DocumentConversionResult(markdown=value)


def _require_markdown_bound(markdown: str, request: DocumentConversionRequest) -> None:
    if len(markdown.encode("utf-8")) > request.max_markdown_bytes:
        raise DocumentConversionError("document_markdown_too_large")


def _table_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", "<br>")


__all__ = [
    "DuckDuckGoSearchProvider",
    "HtmlScrapeProvider",
    "HttpxWebClient",
    "LocalDocumentConverter",
    "PublicWebPolicy",
    "production_run_bindings",
]
